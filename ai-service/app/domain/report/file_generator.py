import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font,PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

def generate_excel_report(data: dict, comments: dict, period: str,
                        start_date: str, end_date: str) -> bytes:
    wb = Workbook() # 엑셀 파일 생성
    wb.remove(wb.active) # 기본 시트 제거
    
    period_label = "월간" if period == "MONTHLY" else "주간"
    
    _build_summary_sheet(wb,data,comments, period_label, start_date, end_date)
    _build_revenue_sheet(wb, data, comments)
    _build_usage_sheet(wb,data, comments)
    
    buffer = io.BytesIO() # 메모리 임시 저장공간 생성
    wb.save(buffer) # 엑셀을 파일이 아닌 메모리에 저장
    return buffer.getvalue() # 바이트 데이터로 반환

# -- 공용 스타일 헬퍼
def _header_fill(color="1E3A5F"): # 셀 배경색 설정(기본값: 네이비 색상)
    return PatternFill("solid", fgColor=color)

def _font(bold=False, size=11, color="000000"): # 글꼴 설정(굵기,크기,색상)
    return Font(bold=bold, size=size, color=color)

def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _thin_border(): # 셀 4면에 얇은 테두리
    side=Side(style="thin")
    return Border(left=side, right=side, top=side, bottom=side)

def _write_header(ws, row, col, value, bg="1E3A5F"): # 헤더 셀 한번에 꾸미기
    cell = ws.cell(row=row, column=col, value=value)
    cell.fill = _header_fill(bg) # 배경색
    cell.font = _font(bold=True, color="FFFFFF") # 흰색 굵은 글씨
    cell.alignment = _center() # 가운데 정렬
    cell.border = _thin_border() # 테두리

def _write_cell(ws, row, col, value, bold=False, align="left"): # 일반 데이터 셀 꾸미기
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = _font(bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = _thin_border()
    return cell

# -- 시트 1: 요약
def _build_summary_sheet(wb,data,comments,period_label,start_date,end_date):
    ws=wb.create_sheet("요약")
    ws.column_dimensions["A"].width=22 # A열 너비
    ws.column_dimensions["B"].width=35 # B열 너비
    
    # 제목
    ws.merge_cells("A1:B1") # A1~B1 셀 합치기
    title = ws["A1"]
    title.value = f"주차장 운영 {period_label} 보고서"
    title.font = _font(bold=True, size=16, color="1E3A5F")
    title.alignment = _center()
    ws.row_dimensions[1].height = 40 # 1행 높이
    
    # 기간
    ws.merge_cells("A2:B2")
    period_cell = ws["A2"]
    period_cell.value = f"기간: {start_date} ~ {end_date}"
    period_cell.font = _font(size=11, color="555555")
    period_cell.alignment = _center()
    
    # 생성일
    ws.merge_cells("A3:B3")
    date_cell = ws["A3"]
    date_cell.value = f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    date_cell.font = _font(size=10,color="888888")
    date_cell.alignment = _center()
    
    ws.append([]) # 빈행
    
    # 현황 헤더
    _write_header(ws, 5, 1, "항목")
    _write_header(ws, 5, 2, "현황")
    
    rows = [
        ("입주 세대수", f"{data.get('household_count','-')} / {data.get('total_household_count','-')}세대"),
        ("등록 차량", f"{data.get('vehicle_count','-')}대"),
        ("주차공간 점유", f"{data.get('occupied_space','-')} / {data.get('total_space','-')}칸"),
        ("주차요금 수익",  f"₩{int(data.get('total_revenue', 0)):,}"),
        ("순 매출액",     f"₩{int(data.get('revenue_total', 0)):,}"),
        ("전월 대비 매출", f"{data.get('revenue_change_percent', '-')}%"),
        ("총 사용량",     f"{data.get('usage_total_count', '-')} 건"),
        ("전월 대비 사용량", f"{data.get('usage_change_percent', '-')}%"),
    ]
    for i, (label,value) in enumerate(rows, start=6): # 6행부터 시작해서 순서대로 작성
        _write_cell(ws,i,1,label,bold=True) # A열: 항목명
        _write_cell(ws,i,2,value,align="center") # B열: 값
    
    # LLM 요약 코멘트
    ws.append([])
    comment_row = 6+len(rows) + 1
    _write_header(ws,comment_row,1,"AI 분석 요약",bg="2E7D32")
    ws.merge_cells(f"A{comment_row}:B{comment_row}")
    
    text_row =comment_row+1
    ws.merge_cells(f"A{text_row}:B{text_row}")
    cell = ws.cell(row=text_row,column=1,value=comments.get("summary_comment",""))
    cell.alignment=Alignment(horizontal="left",vertical="top",wrap_text=True)
    cell.font=_font(size=10)
    ws.row_dimensions[text_row].height=80
    
# -- 시트 2: 매출 현황
def _build_revenue_sheet(wb,data,comments):
    ws=wb.create_sheet("매출 현황")
    ws.column_dimensions["A"].width=18
    ws.column_dimensions["B"].width=20
    
    _write_header(ws,1,1,"매출 현황")
    ws.merge_cells("A1:B1")
    
    rows=[
        ("순 매출액", f"₩{int(data.get('revenue_total', 0)):,}"),
        ("전월 대비", f"{data.get('revenue_change_percent', '-')}%")
    ]
    for i,(label,value) in enumerate(rows,start=2):
        _write_cell(ws,i,1,label,bold=True)
        _write_cell(ws,i,2,value,align="center")
    
    # 월별 매출 추이 테이블
    monthly = data.get("revenue_monthly",[])
    if monthly:
        ws.append([])
        header_row=2+len(rows)+1
        _write_header(ws,header_row,1,"월")
        _write_header(ws,header_row,2,"매출액")
        for j,item in enumerate(monthly,start=header_row+1):
            _write_cell(ws,j,1,item.get("month",""),align="center")
            _write_cell(ws,j,2,f"₩{int(item.get('amount', 0)):,}", align="right")
    
    # LLM 매출 분석 코멘트
    last_row=ws.max_row+2
    _write_header(ws,last_row,1,"AI 매출 분석",bg="2E7D32")
    ws.merge_cells(f"A{last_row}:B{last_row}")
    text_row = last_row+1
    ws.merge_cells(f"A{text_row}:B{text_row}")
    cell = ws.cell(row=text_row, column=1, value=comments.get("revenue_comment", ""))
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    cell.font = _font(size=10)
    ws.row_dimensions[text_row].height = 80

# -- 시트 3: 사용량 현황
def _build_usage_sheet(wb, data, comments):
    ws = wb.create_sheet("사용량 현황")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 20

    _write_header(ws, 1, 1, "사용량 현황")
    ws.merge_cells("A1:B1")

    rows = [
        ("총 사용량",      f"{data.get('usage_total_count', '-')} 건"),
        ("전월 대비",      f"{data.get('usage_change_percent', '-')}%"),
    ]
    for i, (label, value) in enumerate(rows, start=2):
        _write_cell(ws, i, 1, label, bold=True)
        _write_cell(ws, i, 2, value, align="center")

    # LLM 사용량 분석 코멘트
    last_row = ws.max_row + 2
    _write_header(ws, last_row, 1, "AI 사용량 분석", bg="2E7D32")
    ws.merge_cells(f"A{last_row}:B{last_row}")
    text_row = last_row + 1
    ws.merge_cells(f"A{text_row}:B{text_row}")
    cell = ws.cell(row=text_row, column=1, value=comments.get("usage_comment", ""))
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    cell.font = _font(size=10)
    ws.row_dimensions[text_row].height = 80